"""Check twelve-slide orientation revisions against an external frozen baseline.

No scientific analysis, full browser regression or member rehearsal is run.
"""
import argparse
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import fitz
import numpy as np
from lxml import etree
from pptx import Presentation

from add_field_photos_to_presentation import PHOTOS
from revise_orientation_presentation import CHANGES

ALLOWED = {'PACKAGE-SHA256.json', 'PROJECT_STATE.md', 'README.md',
           'docs/UrbanHeat-AlKhor-Pitch.pptx', 'docs/UrbanHeat-AlKhor-Pitch.pdf',
           'docs/Three-Minute-Presentation.md', 'docs/Judge-Questions.md',
           'docs/Team-Review.md', 'docs/Organizer-Access-and-Deadline.md',
           'docs/Submission-Draft.md', 'docs/Submission-Checklist.md',
           'docs/Delivery-Guide.md', 'src/build_delivery.py'}
REGIONS = {2: [(48, 123, 914, 200), (48, 207, 914, 284), (48, 291, 914, 368)],
           3: [(48, 477, 914, 529)], 4: [(48, 348, 914, 426)],
           10: [(48, 207, 914, 284)],
           11: [(48, 250, 914, 304), (48, 378, 914, 455)]}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def compact_spoken(text):
    return [re.search(r'## Slide ' + str(n) + r' — [^\n]+\n\n(.*?)(?=\n## |\Z)',
                      text, re.S).group(1).strip() for n in range(1, 13)]


def verify(root, stage):
    baseline = stage / 'baseline'
    pdf_name = 'docs/UrbanHeat-AlKhor-Pitch.pdf'
    ppt_name = 'docs/UrbanHeat-AlKhor-Pitch.pptx'
    old_ppt = Presentation(baseline / ppt_name)
    new_ppt = Presentation(root / ppt_name)
    assert len(old_ppt.slides) == len(new_ppt.slides) == 12
    assert (old_ppt.slide_width, old_ppt.slide_height) == (new_ppt.slide_width, new_ppt.slide_height)
    allowed_parts = {f'ppt/slides/slide{n}.xml' for n, _ in CHANGES}
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    nodes_changed = 0
    with zipfile.ZipFile(baseline / ppt_name) as old, zipfile.ZipFile(root / ppt_name) as new:
        assert old.testzip() is new.testzip() is None
        assert old.namelist() == new.namelist()
        changed = {n for n in old.namelist() if old.read(n) != new.read(n)}
        assert changed == allowed_parts
        for name in changed:
            a = etree.fromstring(old.read(name)); b = etree.fromstring(new.read(name))
            aa = a.findall('.//a:t', ns); bb = b.findall('.//a:t', ns)
            assert len(aa) == len(bb)
            for before, after in zip(aa, bb):
                nodes_changed += before.text != after.text
                before.text = after.text = ''
            assert etree.tostring(a, method='c14n') == etree.tostring(b, method='c14n')
        assert nodes_changed == len(CHANGES) == 8
        assert old.read('ppt/slides/slide6.xml') == new.read('ppt/slides/slide6.xml')
        assert old.read('ppt/slides/slide8.xml') == new.read('ppt/slides/slide8.xml')
    photo_checks = []
    pictures = [s for s in new_ppt.slides[5].shapes if s.shape_type == 13]
    assert len(pictures) == 3
    expected_photos = set()
    for picture, (filename, case, _) in zip(pictures, PHOTOS):
        original = (root / 'app/assets/site-photos' / filename).read_bytes()
        assert picture.image.blob == original
        expected_photos.add(sha(original))
        photo_checks.append({'case': case, 'filename': filename, 'sha256': sha(original),
                             'pptx_original_bytes': True, 'pdf_original_bytes': True})
    normalize = lambda t: re.sub(r'\s+', '', t)
    text_checks = span_checks = 0
    pixels = []
    with fitz.open(baseline / pdf_name) as old, fitz.open(root / pdf_name) as new:
        assert len(old) == len(new) == 12
        assert expected_photos <= {sha(new.extract_image(i[0])['image']) for i in new[5].get_images()}
        for number, (a, b, slide) in enumerate(zip(old, new, new_ppt.slides), 1):
            assert a.rect == b.rect
            for shape in slide.shapes:
                frames = [shape.text_frame] if shape.has_text_frame else []
                if shape.has_table:
                    frames += [c.text_frame for row in shape.table.rows for c in row.cells]
                for frame in frames:
                    for para in frame.paragraphs:
                        if para.text.strip():
                            assert normalize(para.text) in normalize(b.get_text()), (number, para.text)
                            text_checks += 1
            for block in b.get_text('dict', flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        assert b.rect.contains(fitz.Rect(span['bbox'])), (number, span)
                        assert '\ufffd' not in span['text']
                        span_checks += 1
            pa = a.get_pixmap(alpha=False); pb = b.get_pixmap(alpha=False)
            aa = np.frombuffer(pa.samples, dtype=np.uint8).reshape(pa.height, pa.width, pa.n)
            bb = np.frombuffer(pb.samples, dtype=np.uint8).reshape(pb.height, pb.width, pb.n)
            assert aa.shape == bb.shape
            if number not in REGIONS:
                assert np.array_equal(aa, bb), ('untouched page', number)
                pixels.append({'slide': number, 'unchanged_rendered_pixels': 'identical'})
            else:
                mask = np.ones(aa.shape[:2], dtype=bool)
                for x0, y0, x1, y1 in REGIONS[number]:
                    mask[y0:y1, x0:x1] = False
                assert np.array_equal(aa[mask], bb[mask]), ('outside reviewed paragraphs', number)
                pixels.append({'slide': number, 'pixels_outside_reviewed_regions': 'identical'})
    tree = json.loads((stage / 'baseline-public-tree.json').read_text())
    protected = []
    for item in tree['tree']:
        if item['type'] != 'blob' or item['path'] in ALLOWED:
            continue
        data = (root / item['path']).read_bytes()
        assert hashlib.sha1(b'blob ' + str(len(data)).encode() + b'\0' + data).hexdigest() == item['sha'], item['path']
        protected.append(item['path'])
    inputs = json.loads((root / 'data/sample_input/SHA256.json').read_text())
    tables = json.loads((root / 'metadata/original-results-sha256.json').read_text())
    assert len(inputs) == 35 and len(tables) == 7
    for name, expected in inputs.items():
        assert sha((root / 'data/sample_input' / name).read_bytes()) == expected
    for name, expected in tables.items():
        assert sha((root / name).read_bytes()) == expected
    receipt = json.loads((root / 'metadata/site-media-receipt-2026-10-08.json').read_text())
    for item in receipt['items']:
        assert sha((root / item['public_asset_path']).read_bytes()) == item['sha256']
    compact = (root / 'docs/Three-Minute-Presentation.md').read_text()
    baseline_compact = (baseline / 'docs/Three-Minute-Presentation.md').read_text()
    assert compact_spoken(compact) == compact_spoken(baseline_compact)
    assert sum(len(s.split()) for s in compact_spoken(compact)) == 394
    assert 'not the reported live-pitch limit' in compact
    live = (root / 'docs/Live-Pitch-Up-to-Ten-Minutes.md').read_text()
    timing = json.loads((root / 'metadata/live-pitch-timing-2026-10-09.json').read_text())
    non_demo = [re.search(r'## Slide ' + str(n) + r' — .*?\n\n\*\*Say\*\*\n\n(.*?)\n\n\*\*Cue:',
                         live, re.S).group(1) for n in range(1, 13)]
    demo = re.search(r'\*\*Demo speech\*\*\n\n(.*?)\n\n\*\*Demo actions:', live, re.S).group(1)
    assert sum(len(s.split()) for s in non_demo) == timing['non_demo_words'] == 821
    assert len(demo.split()) == timing['demo_words'] == 74
    assert timing['word_count'] == 895
    assert timing['planned_total_seconds'] == 500 < timing['organizer_reported_pitch_max_seconds'] == 600
    assert timing['separate_question_budget_seconds'] == 300
    assert timing['actual_pitch_rehearsal_seconds'] is timing['actual_qa_rehearsal_seconds'] is None
    assert timing['planning_pause_transition_margin_seconds'] > 0
    assert timing['sections'][0]['start_seconds'] == 0 and timing['sections'][-1]['end_seconds'] == 500
    for a, b in zip(timing['sections'], timing['sections'][1:]):
        assert a['end_seconds'] == b['start_seconds']
    qa = (root / 'docs/Five-Minute-Judge-QA.md').read_text()
    budgets = re.findall(r'\| (\d+):(\d+)–(\d+):(\d+) \|', qa)
    assert len(budgets) == 6
    end = 0
    for values in budgets:
        a, b, c, d = map(int, values)
        assert a * 60 + b == end
        end = c * 60 + d
    assert end == 300
    capture = json.loads((stage / 'demo-capture-verification.json').read_text())
    assert capture['page_errors'] == [] and capture['offline'] is True
    assert capture['source_preview_sha256'] == sha((root / 'preview/UrbanHeat-AlKhor-Preview.html').read_bytes())
    fallback = root / 'docs/Demo-Fallback-Screenshots.pdf'
    assert sha(fallback.read_bytes()) == capture['fallback_pdf']['sha256']
    with fitz.open(fallback) as pdf:
        assert len(pdf) == 2
        embedded = set()
        for page in pdf:
            assert page.get_images()
            assert 'Screenshot, not live success' in page.get_text()
            embedded.update(sha(pdf.extract_image(image[0])['image']) for image in page.get_images())
        assert {item['sha256'] for item in capture['screenshots'].values()} <= embedded
    # Check explicit relative file links in all changed/new user-facing prose.
    docs = ['README.md', 'docs/Three-Minute-Presentation.md', 'docs/Judge-Questions.md',
            'docs/Team-Review.md', 'docs/Organizer-Access-and-Deadline.md',
            'docs/Submission-Draft.md', 'docs/Submission-Checklist.md', 'docs/Delivery-Guide.md',
            'docs/Orientation-Requirements-and-Handoff.md', 'docs/Live-Pitch-Up-to-Ten-Minutes.md',
            'docs/Five-Minute-Judge-QA.md', 'docs/Live-Demo-Runbook.md']
    links = []
    for name in docs:
        p = root / name
        for url in re.findall(r'\]\(([^\s)]+)\)', p.read_text()):
            if '://' in url or url.startswith('#'):
                continue
            path = (p.parent / url.split('#')[0]).resolve()
            assert path.exists(), (name, url)
            links.append({'file': name, 'target': url, 'exists': True})
    external = root / 'docs/Orientation-Requirements-External-Review-2026-10-08.md'
    original = Path('/workspace/attachments/45fec429-4e75-4c6d-9f96-e8c8989ed86d/UrbanHeat-Orientation-Requirements-Review-2026-10-08.md')
    assert external.read_bytes() == original.read_bytes()
    artifacts = [pdf_name, ppt_name, 'docs/Live-Pitch-Up-to-Ten-Minutes.md',
                 'docs/Five-Minute-Judge-QA.md', 'docs/Live-Demo-Runbook.md',
                 'docs/Demo-Fallback-Screenshots.pdf', 'docs/Three-Minute-Presentation.md',
                 'docs/Judge-Questions.md', 'metadata/live-pitch-timing-2026-10-09.json']
    report = {
        'checked_at_utc': datetime.now(timezone.utc).isoformat(), 'user_facing_date': '2026-10-09',
        'user_facing_timezone': 'Asia/Qatar',
        'status': 'orientation_revision_artifact_checks_passed_actual_member_rehearsal_and_native_powerpoint_pending',
        'baseline_public_commit': '8abc67fd2027d6480aac93dbeba690c57620ee0e',
        'baseline_public_files': 260,
        'supplied_review_bytes_preserved': True, 'supplied_review_sha256': sha(external.read_bytes()),
        'underlying_transcript_recordings_live_guide_independently_checked': False,
        'pdf_pages': 12, 'pptx_slides': 12, 'changed_slide_text_parts': [2, 3, 4, 10, 11],
        'text_nodes_changed': nodes_changed,
        'pptx_non_text_xml_geometry_fonts_theme_media_unrelated_parts': 'exact',
        'slide_6_photos_captions_provenance_and_slide_8_table_xml': 'exact unchanged',
        'photographs': photo_checks, 'pdf_pptx_text_matches': text_checks,
        'pdf_text_span_boundary_checks': span_checks, 'pdf_pixel_checks': pixels,
        'visual_review': 'Revised PDF slides 2/3/4/10/11 and retained photo slide 6 inspected at presentation scale; no obvious clipping/overlap. Real screenshot fallback inspected separately.',
        'native_desktop_powerpoint_opened': False,
        'native_render_limit': 'PowerPoint/LibreOffice unavailable; PPTX structural and PDF visual checks do not certify actual-computer font/render acceptance.',
        'pdf_method': 'Retain original pages; targeted existing Noto Sans regular-size/colour/baseline text replacement. No native PowerPoint export claim.',
        'live_pitch_timing': timing, 'compact_spoken_words': 394, 'compact_spoken_text_changed': False,
        'qa_mock_budget_seconds': 300, 'qa_prepared_rounds': 6,
        'demo_capture': capture, 'full_browser_regression_rerun': False,
        'existing_browser_proof': 'results/site-media-capture-verification-2026-10-08.json matches exact unchanged app/preview',
        'relative_links_checked': links,
        'protected_prior_public_files': len(protected), 'protected_files': protected,
        'original_input_hashes_verified': 35, 'original_analytical_csv_hashes_verified': 7,
        'original_received_media_hashes_verified': len(receipt['items']),
        'analysis_rerun': False, 'new_model_or_satellite_data': False, 'ranking_changed': False,
        'primary_dw_confidence': .60, 'primary_dw_common_coverage_percent': 30.938509530941777,
        'unchanged_coverage_gate_percent': 50, 'primary_gate': 'failed',
        'Unknown_excluded_common_percent': 69.06149046905822,
        'independent_scientific_and_member_reviews': 'pending until personally confirmed',
        'actual_rehearsal_invitation_slot_attendance': 'unconfirmed',
        'internal_preparation_target': '2026-10-10; team buffer, not organizer cutoff',
        'exact_organizer_cutoff_live_account_guide': 'unresolved/unverified; conflicting reported sources retained',
        'screen_recording': 'optional according to supplied review; not produced; field video separate',
        'organizer_contact_form_changes_submission': 'not performed',
        'files': {n: {'bytes': (root / n).stat().st_size, 'sha256': sha((root / n).read_bytes())}
                  for n in artifacts},
    }
    (root / 'results/orientation-revision-verification-2026-10-09.json').write_text(json.dumps(report, indent=2) + '\n')
    return {k: report[k] for k in ['status', 'pdf_pages', 'pptx_slides', 'text_nodes_changed',
            'pdf_pptx_text_matches', 'protected_prior_public_files', 'original_input_hashes_verified',
            'original_analytical_csv_hashes_verified', 'compact_spoken_words']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--baseline-stage', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(verify(args.project.resolve(), args.baseline_stage.resolve()), indent=2))
