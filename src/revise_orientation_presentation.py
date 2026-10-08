"""Apply the orientation narrative to a preserved twelve-slide presentation.

Only reviewed text nodes change; all original media, numeric tables and unrelated
PPTX parts remain exact. PDF changes use existing fonts and original baselines.
This does not run analysis or claim native desktop PowerPoint rendering.
"""
import argparse
import zipfile
from pathlib import Path

import fitz
from lxml import etree
from pptx import Presentation

from add_field_photos_to_presentation import BACKGROUND, REGULAR, normalized, wrapped

CHANGES = {
    (2, 1): 'Proposed user: a municipal planning team choosing where to inspect first.',
    (2, 2): 'Shortlist cells, then inspect use, shade, access, ownership and feasibility before choosing an intervention.',
    (2, 3): 'Proposed value: organise evidence and inspection decisions. No adoption or measured benefit is established.',
    (3, 4): '300 m September values are not stop measurements. October field photos show context, not measured use or shade.',
    (4, 4): 'Field photos: 8 October 2026; thermal data: September. No hyperspectral, 813 or proprietary VHR data.',
    (10, 2): 'Proposed pilot: planner feedback and field shade/use/access checks; evaluate shortlist usefulness and feasibility.',
    (11, 3): 'Mohammed Almarri: rehearse the live pitch (up to 10 min) and 5 min Q&A.',
    (11, 5): 'Ali Alkubaisi: rehearse the offline demo and prepare five-minute judge Q&A.',
}


def revise(root, baseline):
    ppt_name = 'docs/UrbanHeat-AlKhor-Pitch.pptx'
    pdf_name = 'docs/UrbanHeat-AlKhor-Pitch.pdf'
    deck = Presentation(baseline / ppt_name)
    assert len(deck.slides) == 12
    changed_parts = {}
    namespace = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
                 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with zipfile.ZipFile(baseline / ppt_name) as source:
        for number in sorted({n for n, _ in CHANGES}):
            name = f'ppt/slides/slide{number}.xml'
            document = etree.fromstring(source.read(name))
            tree = document.find('p:cSld/p:spTree', namespace)
            for (slide, index), text in CHANGES.items():
                if slide != number:
                    continue
                nodes = tree[index + 2].findall('.//a:t', namespace)
                assert len(nodes) == 1
                assert nodes[0].text == deck.slides[number - 1].shapes[index].text
                nodes[0].text = text
            changed_parts[name] = etree.tostring(document, xml_declaration=True,
                                                 encoding='UTF-8', standalone=True)
        with zipfile.ZipFile(root / ppt_name, 'w', zipfile.ZIP_DEFLATED) as target:
            for name in source.namelist():
                target.writestr(name, changed_parts.get(name, source.read(name)))
    font = fitz.Font(fontfile=str(REGULAR))
    with fitz.open(baseline / pdf_name) as pdf:
        for (number, index), text in CHANGES.items():
            page = pdf[number - 1]
            shape = deck.slides[number - 1].shapes[index]
            old_text = shape.text
            blocks = page.get_text('dict', flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)['blocks']
            found = [b for b in blocks if normalized(''.join(s['text'] for line in b.get('lines', [])
                                                            for s in line['spans'])) == normalized(old_text)]
            assert len(found) == 1, (number, old_text)
            block = found[0]
            first = block['lines'][0]['spans'][0]
            x, y = first['origin']
            size = first['size']
            leading = block['lines'][1]['spans'][0]['origin'][1] - y if len(block['lines']) > 1 else size * 1.2
            lines = wrapped(text, font, size, shape.width / 12700 - 14.4)
            bottom = (shape.top + shape.height) / 12700
            assert y + (len(lines) - 1) * leading - font.descender * size <= bottom + 1, (number, lines)
            bbox = fitz.Rect(block['bbox']) + (-1, -1, 1, 1)
            page.add_redact_annot(bbox, fill=BACKGROUND)
            page.apply_redactions(images=0, graphics=0)
            # Redaction can remove an unused font resource from this page.
            # Reattach the same Noto Sans font; garbage=4 merges duplicates.
            name = 'orientation-regular'
            color = tuple(((first['color'] >> shift) & 255) / 255 for shift in (16, 8, 0))
            for i, line in enumerate(lines):
                page.insert_text((x, y + i * leading), line, fontsize=size, fontname=name,
                                 fontfile=str(REGULAR), color=color)
        pdf.save(root / pdf_name, garbage=4, deflate=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--baseline', type=Path, required=True)
    args = parser.parse_args()
    revise(args.project.resolve(), args.baseline.resolve())
